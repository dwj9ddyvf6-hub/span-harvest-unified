from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "RESEARCH_STUDY.md"
TARGET = ROOT / "SPAN_HARVEST_UNIFIED_DEEP_RESEARCH_STUDY.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    relationship = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def clean_heading(value: str) -> str:
    value = re.sub(r"[`*_]", "", value)
    value = re.sub(r"[^0-9A-Za-zÀ-ÿ ]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def normalise_text(value: str) -> str:
    value = value.replace("\\bar F", "average fill")
    value = value.replace("\\mid", "given")
    value = value.replace("\\times", "×")
    value = value.replace("\\ge", "≥")
    value = value.replace("\\le", "≤")
    value = value.replace("−", "-")
    return value


INLINE = re.compile(r"(\[[^\]]+\]\([^\)]+\)|`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def add_inline(paragraph, text: str) -> None:
    text = normalise_text(text)
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("["):
            link = re.match(r"\[([^\]]+)\]\(([^\)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
            else:
                paragraph.add_run(token)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token[1:-1]).italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_body_paragraph(document, text: str, style: str = "Body Text"):
    paragraph = document.add_paragraph(style=style)
    add_inline(paragraph, text)
    return paragraph


def remove_paragraph_border(style) -> None:
    """Remove inherited Word title/header rules that render as decorative lines."""
    properties = style._element.get_or_add_pPr()
    border = properties.find(qn("w:pBdr"))
    if border is not None:
        properties.remove(border)


def add_table(document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for row_index, row in enumerate(rows):
        row_properties = table.rows[row_index]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_properties.append(cant_split)
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, "1F4E79" if row_index == 0 else ("F5F8FB" if row_index % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or col_index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, value.strip())
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = RGBColor(32, 32, 32)
    styles["Body Text"].font.name = "Aptos"
    styles["Body Text"].font.size = Pt(10.5)
    styles["Body Text"].paragraph_format.space_after = Pt(7)
    styles["Body Text"].paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    remove_paragraph_border(styles["Title"])
    styles["Title"].paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "SPAN HARVEST UNIFIED  |  RESEARCH STUDY"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Research artifact  •  17 September 2026")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)

    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []
    first_h1 = True
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.18)
                paragraph.paragraph_format.space_after = Pt(7)
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8.7)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(normalise_text(line))
            continue

        if line.startswith("|") and line.endswith("|"):
            table_lines.append(line)
            continue
        if table_lines:
            rows = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if all(set(cell) <= {"-", ":", " "} for cell in cells):
                    continue
                rows.append(cells)
            add_table(document, rows)
            table_lines = []

        if not line.strip():
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.add_run(clean_heading(line[2:]))
            first_h1 = False
            continue
        if line.startswith("## "):
            document.add_paragraph(clean_heading(line[3:]), style="Heading 1")
            continue
        if line.startswith("### "):
            document.add_paragraph(clean_heading(line[4:]), style="Heading 2")
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, line[2:])
            continue
        if re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph(style="Body Text")
            paragraph.paragraph_format.left_indent = Inches(0.24)
            paragraph.paragraph_format.first_line_indent = Inches(-0.20)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, line)
            continue
        add_body_paragraph(document, line)

    if table_lines:
        rows = []
        for table_line in table_lines:
            cells = [cell.strip() for cell in table_line.strip("|").split("|")]
            if not all(set(cell) <= {"-", ":", " "} for cell in cells):
                rows.append(cells)
        add_table(document, rows)

    core = document.core_properties
    core.title = "Span Harvest Unified Deep Research Study"
    core.subject = "Causal prediction market exit algorithm and validation record"
    core.author = ""
    core.last_modified_by = ""
    core.keywords = ""
    core.comments = ""
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    build()
